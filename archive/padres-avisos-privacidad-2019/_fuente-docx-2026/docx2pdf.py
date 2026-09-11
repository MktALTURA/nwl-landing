"""Convert the NWL 'Aviso de Privacidad' .docx files to PDF.

The sources are plain text: no images, no headers/footers - a title, roman
numeral section headings, body copy and auto-numbered lists.

Two source quirks are handled rather than reproduced:
  * most section headings share a paragraph with the body text that follows
    them, so the run of bold text at the front is split off as the heading;
  * 'página web:' runs straight into the URL with no separating space.
"""
import re, sys, html
import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Paragraph

BODY   = ParagraphStyle('body', fontName='Helvetica', fontSize=10, leading=14.5,
                        alignment=TA_JUSTIFY, spaceAfter=7)
TITLE  = ParagraphStyle('title', parent=BODY, fontSize=15, leading=20,
                        alignment=TA_CENTER, spaceAfter=4)
SUBTIT = ParagraphStyle('sub', parent=BODY, fontSize=11, leading=15,
                        alignment=TA_CENTER, spaceAfter=16)
HEAD   = ParagraphStyle('head', parent=BODY, fontSize=11, leading=15,
                        spaceBefore=12, spaceAfter=6)
ITEM   = ParagraphStyle('item', parent=BODY, leftIndent=1.0*cm,
                        bulletIndent=0.25*cm, spaceAfter=5)


def esc(t):
    return html.escape(t, quote=False)


def markup(runs, force_bold=False):
    out = []
    for r in runs:
        t = esc(r.text)
        if not t:
            continue
        if r.bold or force_bold: t = f'<b>{t}</b>'
        if r.italic:             t = f'<i>{t}</i>'
        if r.underline:          t = f'<u>{t}</u>'
        out.append(t)
    s = re.sub(r'\s+', ' ', ''.join(out)).strip()
    # the URL may sit in its own run, so allow inline tags between the two
    return re.sub(r':((?:<[^>]+>)*)(https?://)', r': \1\2', s)


def eff_bold(run, para_bold):
    """A run's bold=None means 'inherit from the paragraph style'."""
    return para_bold if run.bold is None else bool(run.bold)


def is_list_item(p):
    return p._p.pPr is not None and p._p.pPr.numPr is not None


STALE_MAIL = 'direccion@sanmiguel.nwl.mx'
OPTOUT_CUE = 'manifestar su negativa a través del correo electrónico'


def localize(d, email):
    """Point every contact address at the campus that owns the document.

    The sources were all derived from the San Miguel master, so each one
    carries that campus's address in the ARCO, opt-out and enquiries
    sections, and four of the five leave the opt-out address blank. The
    signature-block blanks ('Nombre del Alumno(a):', 'Fecha:') are form
    fields and are deliberately left alone.
    """
    mails = blanks = 0
    for p in d.paragraphs:
        optout = OPTOUT_CUE in p.text
        for r in p.runs:
            if STALE_MAIL in r.text:
                r.text = r.text.replace(STALE_MAIL, email)
                mails += 1
            if optout and re.search(r'_{3,}', r.text):
                r.text = re.sub(r'_{3,}', email, r.text)
                blanks += 1
    return mails, blanks


def build(src, dest, footer_label, email=None):
    d = docx.Document(src)
    stats = localize(d, email) if email else (0, 0)
    story, counter, seen_title = [], 0, False

    for p in d.paragraphs:
        if not p.text.strip():
            continue
        runs = list(p.runs)
        style = (p.style.name or '').lower()

        if is_list_item(p):
            counter += 1
            story.append(Paragraph(markup(runs), ITEM, bulletText=f'{counter}.'))
            continue
        counter = 0

        if style == 'title':
            story.append(Paragraph(markup(runs, force_bold=True), TITLE))
            seen_title = True
        elif not seen_title:
            story.append(Paragraph(markup(runs, force_bold=True), SUBTIT))
            seen_title = True
        elif style.startswith('heading'):
            # split the leading bold run(s) off as the heading proper
            n = 0
            pb = bool(p.style.font.bold)
            while n < len(runs) and (eff_bold(runs[n], pb) or not runs[n].text.strip()):
                n += 1
            while n and not runs[n-1].text.strip():
                n -= 1          # don't let the heading keep a trailing space run
            head, rest = (runs[:n], runs[n:]) if n else (runs, [])
            story.append(Paragraph(markup(head, force_bold=True), HEAD))
            if rest:
                story.append(Paragraph(markup(rest), BODY))
        else:
            story.append(Paragraph(markup(runs), BODY))

    def footer(canv, doc):
        canv.saveState()
        canv.setFont('Helvetica', 8); canv.setFillGray(0.45)
        canv.drawCentredString(A4[0]/2.0, 1.35*cm,
                               f'{footer_label}  ·  Página {doc.page}')
        canv.restoreState()

    doc = BaseDocTemplate(dest, pagesize=A4,
                          leftMargin=2.54*cm, rightMargin=2.54*cm,
                          topMargin=2.54*cm, bottomMargin=2.3*cm,
                          title='Aviso de Privacidad', author='Colegio Newland S.C.')
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id='n')
    doc.addPageTemplates([PageTemplate(id='p', frames=[frame], onPage=footer)])
    doc.build(story)
    return stats


if __name__ == '__main__':
    print(build(*sys.argv[1:5]))
